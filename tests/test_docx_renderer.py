"""docx_renderer 单元测试：SSIR JSON → docx（doc_1，与 render.pdf 内容等价）。

覆盖：封面/目次/正文结构、表格合并单元格语义（vMerge/gridSpan）、
表格单元格脚注标记与平拍指数上标还原、真实 canonical 解析冒烟。
"""

import tempfile
import unittest
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from leleby_ssir.docx_importer import docx_metadata, docx_to_csm_markdown
from leleby_ssir.docx_renderer import HAN_PT, _DocxBuilder, render_docx_file, render_ssir_docx
from leleby_ssir.service import normalize_csm, parse_csm, round_trip_csm


def _mini_document() -> dict:
    """结构最小但覆盖主要内容的 SSIR 文档（绕过 schema 校验，直接喂渲染器）。"""
    return {
        "documentType": "standard",
        "metadata": {
            "common": {
                "documentIdentifier": "GB/T 12345-2021",
                "title": "测试标准",
                "titleEn": "Test standard",
                "issuer": "测试发布机构",
                "publicationDate": "2021-01-01",
                "effectiveDate": "2021-02-01",
            },
            "standard": {"standardNumber": "GB/T 12345-2021", "replaces": "GB/T 9999-2010"},
        },
        "structuralRoot": {
            "children": [
                {
                    "id": "n-toc", "nodeType": "documentBlock", "level": 1, "title": "目 次",
                    "sortOrder": 0,
                    "contentElements": [
                        {"id": "c-toc-list", "presentationType": "list", "sortOrder": 0,
                         "listItems": [{"id": "li-1", "marker": "-", "text": "1 范围", "sortOrder": 0}]},
                    ],
                },
                {
                    "id": "n-1", "nodeType": "section", "level": 1, "title": "范围", "number": "1",
                    "sortOrder": 1,
                    "contentElements": [
                        {"id": "c-p1", "presentationType": "paragraph",
                         "textContent": "本文件规定了测试要求，见 6.62607015×10⁻³⁴ J·s。",
                         "sortOrder": 0},
                    ],
                },
                {
                    "id": "n-2", "nodeType": "section", "level": 1, "title": "表2 样例",
                    "number": "2", "sortOrder": 2,
                    "contentElements": [
                        {"id": "c-table", "presentationType": "table", "tableRef": "t1", "sortOrder": 0},
                        {"id": "c-note", "presentationType": "note",
                         "textContent": "注：表中数值为示例。", "sortOrder": 1},
                    ],
                },
                {
                    "id": "n-a", "nodeType": "annex", "level": 1,
                    "title": "（资料性） 示例附录", "number": "A",
                    "sortOrder": 3,
                    "contentElements": [
                        {"id": "c-pa", "presentationType": "paragraph",
                         "textContent": "5.3.1平排条号段落。", "sortOrder": 0},
                    ],
                },
            ]
        },
        "tables": [
            {
                "id": "t1", "number": "1", "caption": "测试表", "unit": "毫米",
                "rows": [
                    {"rowIndex": 0, "isHeader": True, "cells": [
                        {"colIndex": 0, "text": "项目"},
                        {"colIndex": 1, "text": "取值", "rowspan": 2},
                        {"colIndex": 2, "text": "备注"},
                    ]},
                    {"rowIndex": 1, "isHeader": False, "cells": [
                        {"colIndex": 0, "text": "1030"},
                        {"colIndex": 2, "text": "匝间绝缘[:^a]"},
                    ]},
                    {"rowIndex": 2, "isHeader": False, "cells": [
                        {"colIndex": 0, "text": "10-2"},
                        {"colIndex": 1, "text": "s−1"},
                        {"colIndex": 2, "text": ""},
                    ]},
                ],
            }
        ],
        "figures": [],
        "formulas": [],
        "unknownContents": [],
        "sourceFiles": [],
    }


class DocxRenderStructureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.output = Path(self.temp.name) / "mini.docx"

    def tearDown(self) -> None:
        self.temp.cleanup()

    def render(self, document: dict) -> list[str]:
        # 结构最小的合成文档不满足完整 schema（validate 路径由
        # DocxRenderGoldenSmokeTests 真实 canonical 覆盖），直接驱动渲染器。
        return _DocxBuilder(document).render(str(self.output))

    def read(self) -> Document:
        return Document(str(self.output))

    def test_cover_toc_body_present(self) -> None:
        warnings = self.render(_mini_document())
        self.assertEqual(warnings, [])
        doc = self.read()
        paragraphs = [p.text for p in doc.paragraphs]
        joined = "\n".join(paragraphs)
        # 封面信息（标准号/名称/机构/日期/代替/英文名）
        for expected in ("GB/T 12345-2021", "测试标准", "Test standard", "测试发布机构",
                         "2021-01-01 发布", "2021-02-01 实施", "代替 GB/T 9999-2010"):
            self.assertIn(expected, joined, expected)
        # 标题（Word 大纲；编号与标题间是 U+3000 全角空格）
        headings = [p.text.replace("\u3000", " ") for p in doc.paragraphs if p.style.name.startswith("Heading")]
        self.assertTrue(any(text == "1 范围" for text in headings), headings)
        self.assertTrue(any("附录 A" in text and "资料性" in text for text in headings), headings)
        # 注释与正文
        self.assertTrue(any("注：表中数值为示例。" in text for text in paragraphs))
        self.assertTrue(any("本文件规定了测试要求" in text for text in paragraphs))

    def test_list_marker_uses_tab_stop_not_space(self) -> None:
        # GBT-B04：列项 marker 与文字之间用制表符 + 显式制表位（文字恰在回行位置），
        # 不用空格——Word 两端对齐会拉伸空格，marker 后空隙忽宽忽窄（PDF 侧同因，
        # 见 pdf_renderer._list_marker_gap）。汉字位按 GB/T 1.1-2020 10.2.2：
        # 第一层次 marker 空两个汉字起排、文字列第四个（回行同）汉字位；第二层次
        # marker 空四个汉字、文字列第六个汉字位（旧实现两级共用 0.85cm）。
        document = _mini_document()
        document["structuralRoot"]["children"].append({
            "id": "n-list", "nodeType": "section", "level": 1, "title": "列项", "number": "9", "sortOrder": 9,
            "contentElements": [
                {"id": "c-list", "presentationType": "list", "sortOrder": 0,
                 "listItems": [
                     {"id": "li-a", "marker": "a）", "text": "海拔不超过1000m；", "sortOrder": 0},
                     {"id": "li-b", "marker": "b）", "text": "当运行地点的海拔超过1000m或运行地点的环境空气温度随海拔升高而下降时，电动机温升限值的修正按GB 755的规定。", "sortOrder": 1},
                     {"id": "li-1", "marker": "1）", "text": "左向(含左上、左下)，图形符号应位于右侧。", "sortOrder": 2},
                 ]},
            ],
        })
        self.render(document)
        doc = self.read()
        paragraphs = [p for p in doc.paragraphs if p.style.name == "SSIR List Item"]
        self.assertEqual(len(paragraphs), 3)
        self.assertEqual(paragraphs[0].text, "a）\t海拔不超过1000m；")
        self.assertIn("<w:tab/>", paragraphs[0]._p.xml)
        self.assertIn("<w:tabs>", paragraphs[0]._p.xml)
        self.assertAlmostEqual(paragraphs[0].paragraph_format.left_indent.pt, 4 * HAN_PT, places=1)
        self.assertAlmostEqual(paragraphs[0].paragraph_format.first_line_indent.pt, -2 * HAN_PT, places=1)
        # 第二层次（数字编号 1)）：文字列 6 汉字、marker 起排 4 汉字。
        self.assertAlmostEqual(paragraphs[2].paragraph_format.left_indent.pt, 6 * HAN_PT, places=1)
        self.assertAlmostEqual(paragraphs[2].paragraph_format.first_line_indent.pt, -2 * HAN_PT, places=1)

    def test_table_merges_and_superscripts(self) -> None:
        self.render(_mini_document())
        doc = self.read()
        table = doc.tables[0]
        # 表题注与单位行
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any(t == "表1 测试表" for t in texts))
        self.assertTrue(any(t == "单位为毫米" for t in texts))
        # 合并：取值 col1 跨两行；物理行1 是 [1030, vMerge continue, 匝间绝缘a] 三个 tc
        self.assertEqual(table.rows[0].cells[1].text.strip(), "取值")
        self.assertEqual(table.rows[1].cells[1].text.strip(), "取值")  # 虚拟网格延续
        tr1_tcs = table.rows[1]._tr.findall(qn("w:tc"))
        self.assertEqual(len(tr1_tcs), 3)
        vmerge = tr1_tcs[1].find(qn("w:tcPr")).find(qn("w:vMerge"))
        self.assertIsNotNone(vmerge)
        # 上标还原：1030 → 10 + sup(30)；s−1 → s + sup(−1)
        sup_runs: list[str] = []
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            if run.font.superscript:
                                sup_runs.append(run.text)
        self.assertIn("30", sup_runs)
        self.assertIn("−1", sup_runs)
        # 表注引用点（显式 [:^a] 标记）：匝间绝缘[:^a] 的 a 应为上标
        self.assertTrue(any(run.font.superscript and run.text == "a" for run in _all_runs(doc)))
        # 表头居中、数据行居左（2026-09-07 用户裁定）
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        header_align = table.rows[0].cells[0].paragraphs[0].alignment
        body_align = table.rows[1].cells[0].paragraphs[0].alignment
        self.assertEqual(header_align, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(body_align, WD_ALIGN_PARAGRAPH.LEFT)

    def test_plain_text_units_not_marked_superscript(self) -> None:
        document = _mini_document()
        cell_texts = ["长度m", "规格mm", "a)中所述"]
        document["tables"][0]["rows"][2]["cells"] = [
            {"colIndex": 0, "text": cell_texts[0]},
            {"colIndex": 1, "text": cell_texts[1]},
            {"colIndex": 2, "text": cell_texts[2]},
        ]
        self.render(document)
        doc = self.read()
        sup = set()
        for run in _all_runs(doc):
            if run.font.superscript:
                sup.add(run.text)
        # 单位 m/mm 与列项引用 a) 不得被误判为脚注上标
        # （表内仍保留真实脚注标记 a 与指数 30 的上标，见上行数据）
        self.assertNotIn("m", sup)
        self.assertNotIn("mm", sup)
        self.assertNotIn("n", sup)


def _all_runs(doc) -> list:
    runs = []
    for p in doc.paragraphs:
        runs.extend(p.runs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    runs.extend(p.runs)
    return runs


def _table_cells(doc) -> list:
    return [cell for table in doc.tables for row in table.rows for cell in row.cells]


@unittest.skipUnless(
    Path("corpus/golden/csm/T_ZZB_1064-2019.canonical.md").is_file(),
    "corpus/golden/csm fixture unavailable",
)
class DocxRenderGoldenSmokeTests(unittest.TestCase):
    """真实语料 canonical → SSIR → docx 冒烟（schema 全量校验路径）。"""

    def test_real_standard_render_smoke(self) -> None:
        fixture = "corpus/golden/csm/T_ZZB_1064-2019.canonical.md"
        with tempfile.TemporaryDirectory() as temp:
            ssir = parse_csm(fixture)
            output = Path(temp) / "doc.docx"
            warnings = render_ssir_docx(ssir, str(output), input_file=fixture)
            self.assertEqual(warnings, [])
            doc = Document(str(output))
            self.assertGreater(len(doc.paragraphs), 20)
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            self.assertTrue(headings)
            # docx 应可被 python-docx 重新打开且标题含章号
            self.assertTrue(any(h.startswith("1") for h in headings if h[:1].isdigit()))


class DocxImportRoundtripTests(unittest.TestCase):
    """docx（doc_1 渲染产物）→ canonical 级 CSM 导入，闭环自洽。"""

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp.name)
        self.docx_path = self.temp_path / "GB_T_12345-2021.docx"
        _DocxBuilder(_mini_document()).render(str(self.docx_path))

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_metadata_restored_from_hidden_payload(self) -> None:
        metadata = docx_metadata(self.docx_path)
        self.assertEqual(metadata["title"], "测试标准")
        self.assertEqual(metadata["standard-number"], "GB/T 12345-2021")
        self.assertEqual(metadata["issuer"], "测试发布机构")
        self.assertEqual(metadata["publication-date"], "2021-01-01")
        self.assertEqual(metadata["document-type"], "standard")

    def test_import_markdown_structure(self) -> None:
        fm, md, warnings = docx_to_csm_markdown(self.docx_path, assets_root=self.temp_path)
        self.assertEqual(warnings, [])
        self.assertEqual(fm["title"], "测试标准")
        # 标题：docx Heading1（章）→ md ##（与 render.md 层级一致）
        self.assertIn("# 测试标准", md)
        self.assertIn("## 1 范围", md)
        self.assertIn("## 目 次", md)
        self.assertIn("## 2 表2 样例", md)
        self.assertIn("## 附录 A （资料性） 示例附录", md)
        # 表格：指令 + 加粗题注 + pipe + merge 坐标（row 0-based / column 1-based）
        self.assertIn("<!-- ssir:table id=\"docx-import-t001\" header-rows=\"1\" -->", md)
        self.assertIn("**表1 测试表**", md)
        self.assertIn("| 项目 | 取值 | 备注 |", md)
        self.assertIn('row="0" column="2" rowspan="2" colspan="1"', md)
        # 上标字形还原：s−1 → s⁻¹（表格单元），脚注标记字母原样保留
        self.assertIn("s⁻¹", md)
        self.assertIn("匝间绝缘a", md)

    def test_import_then_normalize_roundtrip(self) -> None:
        fm, md, _ = docx_to_csm_markdown(self.docx_path, assets_root=self.temp_path)
        raw = self.temp_path / "imported.md"
        raw.write_text(md, encoding="utf-8")
        canonical = self.temp_path / "imported.canonical.md"
        _, report = normalize_csm(raw, canonical)
        self.assertIsNotNone(report)
        render_md = self.temp_path / "imported.render.md"
        _ssir, _verify, rt = round_trip_csm(canonical, render_md)
        self.assertTrue(rt.passed)


class DocxRenderFileApiTests(unittest.TestCase):
    def test_render_docx_file_requires_existing_json(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "missing.ssir.json"
            with self.assertRaises(OSError):
                render_docx_file(str(source), str(Path(temp) / "x.docx"))


class DocxFormulaNumberLineTests(unittest.TestCase):
    """GBT-X06（GB/T 1.1-2020 9.9.2、10.4.3）：公式编号右端对齐、与公式以“……”连接。

    Word 侧用制表位实现：居中制表位（按「公式 + 两个汉字间隔 + 引导线」整段居中
    补偿）+ 右端制表位；省略号个数按版心宽度计算（与 PDF 同规）。docx 回灌
    （docx_importer）时制表符/引导线作为版式噪声剥离，编号还原为「式(N)」行。
    """

    def _render(self, directory: str) -> Path:
        from PIL import Image as PILImage

        root = Path(directory)
        asset_dir = root / "assets" / "images"
        asset_dir.mkdir(parents=True, exist_ok=True)
        PILImage.new("RGB", (443, 74), (255, 255, 255)).save(asset_dir / "eq1.jpg")
        ssir = _mini_document()
        ssir["formulas"] = [
            {"id": "fm-1", "rawText": "\\Delta t = x", "latex": "\\Delta t = x",
             "number": "1", "assetRef": "assets/images/eq1.jpg"},
        ]
        ssir["structuralRoot"]["children"][1]["contentElements"].append(
            {"id": "c-f1", "presentationType": "formula", "formulaRef": "fm-1", "sortOrder": 1}
        )
        ssir_path = root / "doc.ssir.json"
        ssir_path.write_text("{}", encoding="utf-8")
        output = root / "doc.docx"
        # 结构最小的合成文档不满足完整 schema（validate 路径由真实 canonical 冒烟覆盖）。
        warnings = _DocxBuilder(ssir, input_file=str(ssir_path)).render(str(output))
        self.assertEqual(warnings, [])
        return output

    def test_number_uses_center_and_right_tab_stops_with_leader_dots(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            output = self._render(temp)
            document = Document(str(output))
            paragraph = next(
                p for p in document.paragraphs if "…" in p.text and p.text.rstrip().endswith("(1)")
            )
            stops = list(paragraph.paragraph_format.tab_stops)
            self.assertEqual(len(stops), 2)
            self.assertEqual(str(stops[0].alignment), "CENTER (1)")
            self.assertEqual(str(stops[1].alignment), "RIGHT (2)")
            self.assertAlmostEqual(stops[1].position.cm, 16.0, places=2)
            # 公式图被收到「留得下引导线」的宽度（只缩小不放大）并居中于制表位。
            self.assertIn("<w:drawing>", paragraph._p.xml)
            self.assertGreaterEqual(paragraph.text.count("…"), 4)
            self.assertLessEqual(len(paragraph.text), 30)

    def test_import_strips_leader_and_restores_number_line(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            output = self._render(temp)
            _metadata, markdown, warnings = docx_to_csm_markdown(output, assets_root=Path(temp))
            self.assertEqual(warnings, [])
            self.assertIn("式(1)", markdown)
            self.assertNotIn("…", markdown)


class DocxNoteExampleLabelFontTests(unittest.TestCase):
    """注/示例标记黑体（2026-09-11，GB/T 1.1-2020 10.4.4.1/10.4.5、附录 F 表 F.1
    序号 42/44；GBT-B10/GBT-B11 执行侧）。

    与 PDF 侧 ``pdf_renderer._label_markup`` 同一规则：标记（"注："/"注1："/"示例1："）
    的 run 用黑体、内容 run 保持宋体；"注意：…"与无冒号的"示例1示出了…"不是标记。
    """

    @staticmethod
    def _document() -> dict:
        document = _mini_document()
        document["structuralRoot"]["children"].append({
            "id": "n-labels", "nodeType": "section", "level": 1, "title": "标记", "number": "3", "sortOrder": 4,
            "contentElements": [
                {"id": "c-note1", "presentationType": "note", "textContent": "注1：注标记应为黑体。", "sortOrder": 0},
                {"id": "c-ex", "presentationType": "example", "textContent": "示例1：正文示例标记也应加黑。", "sortOrder": 1},
                {"id": "c-neg1", "presentationType": "paragraph", "textContent": "注意：此处的“注意”不是注标记。", "sortOrder": 2},
                {"id": "c-neg2", "presentationType": "paragraph", "textContent": "示例1示出了行内引用形态。", "sortOrder": 3},
                {"id": "c-table", "presentationType": "table", "tableRef": "t-note", "sortOrder": 4},
            ],
        })
        document["tables"].append({
            "id": "t-note", "number": "9", "caption": "表注标记", "unit": "",
            "rows": [
                {"rowIndex": 0, "isHeader": True, "cells": [
                    {"colIndex": 0, "text": "项目"},
                    {"colIndex": 1, "text": "取值"},
                ]},
                {"rowIndex": 1, "isHeader": False, "cells": [
                    {"colIndex": 0, "text": "a"},
                    {"colIndex": 1, "text": "1.0"},
                ]},
                {"rowIndex": 2, "isHeader": False, "cells": [
                    {"colIndex": 0, "text": "注2：表内注标记应为黑体。", "colspan": 2},
                ]},
            ],
        })
        return document

    @staticmethod
    def _run_fonts(run: Any) -> tuple[str | None, bool | None]:
        rpr = run._element.rPr
        east_asian = None
        if rpr is not None:
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                east_asian = fonts.get(qn("w:eastAsia"))
        return east_asian, run.font.bold

    def test_note_and_example_labels_use_hei_and_content_stays_song(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "labels.docx"
            warnings = _DocxBuilder(self._document()).render(str(output))
            self.assertEqual(warnings, [])
            document = Document(str(output))
            paragraphs = {p.text.strip(): p for p in document.paragraphs}
            for text, label in (("注1：注标记应为黑体。", "注1："), ("示例1：正文示例标记也应加黑。", "示例1：")):
                with self.subTest(text=text):
                    paragraph = paragraphs[text]
                    self.assertEqual(paragraph.runs[0].text, label)
                    self.assertEqual(self._run_fonts(paragraph.runs[0]), ("黑体", True))
                    self.assertEqual(self._run_fonts(paragraph.runs[1])[0], "宋体")
                    self.assertFalse(paragraph.runs[1].font.bold)
            # 反例：无冒号引用与"注意"都不加黑。
            for text in ("注意：此处的“注意”不是注标记。", "示例1示出了行内引用形态。"):
                with self.subTest(text=text):
                    paragraph = paragraphs[text]
                    for run in paragraph.runs:
                        self.assertNotEqual(self._run_fonts(run), ("黑体", True))

    def test_table_note_cell_label_uses_hei(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "labels.docx"
            _DocxBuilder(self._document()).render(str(output))
            document = Document(str(output))
            table = next(t for t in document.tables if any("注2：" in c.text for row in t.rows for c in row.cells))
            cell = next(c for row in table.rows for c in row.cells if "注2：" in c.text)
            paragraph = next(p for p in cell.paragraphs if p.text.strip().startswith("注2："))
            self.assertEqual(paragraph.runs[0].text, "注2：")
            self.assertEqual(self._run_fonts(paragraph.runs[0]), ("黑体", True))
            self.assertEqual(self._run_fonts(paragraph.runs[1])[0], "宋体")


if __name__ == "__main__":
    unittest.main()
